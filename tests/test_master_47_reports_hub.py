"""Unit and integration tests for MASTER-47: Advanced Analytics & Reporting Hub in Web Admin."""

from io import BytesIO
from unittest.mock import AsyncMock, MagicMock, patch
import uuid
import pytest
from starlette.testclient import TestClient
import openpyxl

from khatmsaz.modules.authorization.models import AdminRole, AdminRoleGrant
from khatmsaz.modules.identity.models import User
from khatmsaz.modules.khatm.models import Khatm, KhatmStatus, KhatmTemplateType
from khatmsaz.modules.wallet.models import Wallet, WalletInvoice
from khatmsaz.web.app import app


@pytest.fixture
def mock_admin():
    admin = MagicMock(spec=User)
    admin.id = uuid.uuid4()
    admin.display_name = "Super Admin"
    admin.is_active = True
    return admin


@pytest.fixture
def mock_admin_grant(mock_admin):
    grant = MagicMock(spec=AdminRoleGrant)
    grant.user_id = mock_admin.id
    grant.role = AdminRole.ADMIN
    grant.revoked_at = None
    return grant


def test_reports_page_unauthenticated():
    """Unauthenticated access to /reports redirects to login."""
    client = TestClient(app, follow_redirects=False)
    response = client.get("/reports")
    assert response.status_code == 303
    assert response.headers["location"] == "/login"


def test_reports_page_authenticated(mock_admin, mock_admin_grant):
    """Authenticated admin can view /reports with KPI cards and 5 categories."""
    client = TestClient(app)

    async def fake_admin(request, *args, **kwargs):
        return mock_admin, "raw_cookie"

    with patch("khatmsaz.web.app._admin", side_effect=fake_admin):
        response = client.get("/reports")
        assert response.status_code == 200
        html = response.text
        assert "مرکز گزارشات و تحلیل‌های آماری" in html
        assert "کاربران فعال امروز" in html
        assert "ماتریس استقبال ۵ خانواده معنوی" in html
        assert "خروجی‌های استاندارد اکسل" in html
        assert "خانواده معنوی" in html


def test_reports_export_khatms_xlsx(mock_admin):
    """Admin can export all khatms as valid XLSX with RTL enabled."""
    client = TestClient(app)

    fake_khatm = MagicMock(spec=Khatm)
    fake_khatm.id = uuid.uuid4()
    fake_khatm.title = "ختم قرآن صلواتی"
    fake_khatm.niyyat = "تعجیل در فرج"
    fake_khatm.template_type = KhatmTemplateType.QURAN_PAGE
    fake_khatm.status = KhatmStatus.ACTIVE
    fake_khatm.is_commitment = True
    fake_khatm.created_at = None

    fake_user = MagicMock(spec=User)
    fake_user.display_name = "بانی نمونه"
    fake_user.phone_number = "09120000000"

    mock_res = MagicMock()
    mock_res.all.return_value = [(fake_khatm, fake_user, None)]

    mock_part_res = MagicMock()
    mock_part_res.all.return_value = [(fake_khatm.id, 10)]

    mock_session = AsyncMock()
    mock_session.execute.side_effect = [mock_res, mock_part_res]

    fake_scope = MagicMock()
    fake_scope.__aenter__.return_value = mock_session
    fake_scope.__aexit__.return_value = None

    async def fake_admin(request, *args, **kwargs):
        return mock_admin, "raw_cookie"

    with patch("khatmsaz.web.app._admin", side_effect=fake_admin), \
         patch("khatmsaz.web.app.session_scope", return_value=fake_scope):
        response = client.get("/reports/export/khatms.xlsx")
        assert response.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers["content-type"]
        assert 'filename="khatmsaz-khatms.xlsx"' in response.headers["content-disposition"]

        # Parse with openpyxl to verify valid xlsx structure
        wb = openpyxl.load_workbook(BytesIO(response.content))
        assert "ختم‌ها" in wb.sheetnames
        sheet = wb["ختم‌ها"]
        assert sheet.sheet_view.rightToLeft is True
        headers = [cell.value for cell in sheet[1]]
        assert "عنوان ختم" in headers
        assert "نام سازنده" in headers


def test_reports_export_users_xlsx(mock_admin):
    """Admin can export all users as valid XLSX."""
    client = TestClient(app)

    fake_user = MagicMock(spec=User)
    fake_user.id = uuid.uuid4()
    fake_user.display_name = "کاربر تستی"
    fake_user.phone_number = "09350000000"
    fake_user.created_at = None

    fake_wallet = MagicMock(spec=Wallet)
    fake_wallet.balance_toman = 50000
    fake_wallet.reward_credit_toman = 10000

    mock_res = MagicMock()
    mock_res.all.return_value = [(fake_user, None, fake_wallet)]

    mock_part_res = MagicMock()
    mock_part_res.all.return_value = [(fake_user.id, 3)]

    mock_session = AsyncMock()
    mock_session.execute.side_effect = [mock_res, mock_part_res]

    fake_scope = MagicMock()
    fake_scope.__aenter__.return_value = mock_session
    fake_scope.__aexit__.return_value = None

    async def fake_admin(request, *args, **kwargs):
        return mock_admin, "raw_cookie"

    with patch("khatmsaz.web.app._admin", side_effect=fake_admin), \
         patch("khatmsaz.web.app.session_scope", return_value=fake_scope):
        response = client.get("/reports/export/users.xlsx")
        assert response.status_code == 200
        assert 'filename="khatmsaz-users.xlsx"' in response.headers["content-disposition"]

        wb = openpyxl.load_workbook(BytesIO(response.content))
        assert "کاربران" in wb.sheetnames
        sheet = wb["کاربران"]
        assert sheet.sheet_view.rightToLeft is True
        headers = [cell.value for cell in sheet[1]]
        assert "نام نمایشی" in headers
        assert "موجودی کیف پول (تومان)" in headers


def test_reports_export_finances_xlsx(mock_admin):
    """Admin can export all finances/invoices as valid XLSX."""
    client = TestClient(app)

    fake_inv = MagicMock(spec=WalletInvoice)
    fake_inv.id = uuid.uuid4()
    fake_inv.net_amount_toman = 75000
    fake_inv.amount_toman = 75000
    fake_inv.description = "خرید پلن حرفه‌ای"
    fake_inv.status = "PAID"
    fake_inv.gateway_ref = "PAYPING-12345"
    fake_inv.issued_at = None
    fake_inv.paid_at = None

    fake_user = MagicMock(spec=User)
    fake_user.display_name = "خریدار پلن"
    fake_user.phone_number = "09121111111"

    mock_res = MagicMock()
    mock_res.all.return_value = [(fake_inv, fake_user)]

    mock_session = AsyncMock()
    mock_session.execute.return_value = mock_res

    fake_scope = MagicMock()
    fake_scope.__aenter__.return_value = mock_session
    fake_scope.__aexit__.return_value = None

    async def fake_admin(request, *args, **kwargs):
        return mock_admin, "raw_cookie"

    with patch("khatmsaz.web.app._admin", side_effect=fake_admin), \
         patch("khatmsaz.web.app.session_scope", return_value=fake_scope):
        response = client.get("/reports/export/finances.xlsx")
        assert response.status_code == 200
        assert 'filename="khatmsaz-finances.xlsx"' in response.headers["content-disposition"]

        wb = openpyxl.load_workbook(BytesIO(response.content))
        assert "تراکنش‌های مالی" in wb.sheetnames
        sheet = wb["تراکنش‌های مالی"]
        assert sheet.sheet_view.rightToLeft is True
        headers = [cell.value for cell in sheet[1]]
        assert "مبلغ (تومان)" in headers
        assert "کد رهگیری درگاه" in headers


def test_reports_export_unauthenticated():
    """Unauthenticated calls to all export endpoints redirect to /login."""
    client = TestClient(app, follow_redirects=False)
    for path in ("/reports/export/khatms.xlsx", "/reports/export/users.xlsx", "/reports/export/finances.xlsx"):
        resp = client.get(path)
        assert resp.status_code == 303
        assert resp.headers["location"] == "/login"
